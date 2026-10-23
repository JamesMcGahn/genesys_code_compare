from unittest.mock import Mock

import pytest
from openpyxl import load_workbook

from genesys_code_compare.cli import CLI
from genesys_code_compare.file.sheet_service import SheetService
from genesys_code_compare.models.file import SheetData
from genesys_code_compare.file.util import FileUtil
import datetime


@pytest.fixture
def service():
    return SheetService(cli=Mock(spec=CLI))


def test_write_creates_readable_workbook(service, tmp_path):
    output = tmp_path / "reports" / "test.xlsx"
    sheets = [
        SheetData(
            name="Summary",
            headers=["Name", "Count", "Enabled"],
            rows=[
                ["Lunch", 2, True],
                ["Break", 1, False],
            ],
        ),
        SheetData(
            name="Details",
            headers=["Source"],
            rows=[["BU A"]],
        ),
    ]

    result = service.write(output, sheets)

    assert result == output
    assert output.is_file()

    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["Summary", "Details"]

        summary = workbook["Summary"]
        assert list(summary.values) == [
            ("Name", "Count", "Enabled"),
            ("Lunch", 2, True),
            ("Break", 1, False),
        ]
        assert summary["A1"].font.bold is True
        assert summary.freeze_panes == "A2"
        assert summary.auto_filter.ref == "A1:C3"

        assert list(workbook["Details"].values) == [
            ("Source",),
            ("BU A",),
        ]
    finally:
        workbook.close()


def test_write_rejects_empty_sheet_list(service, tmp_path):
    output = tmp_path / "empty.xlsx"

    with pytest.raises(ValueError, match="Provide at least one sheet"):
        service.write(output, [])

    assert not output.exists()


def test_default_output(service, tmp_path, monkeypatch):
    monkeypatch.setattr(
        FileUtil,
        "get_app_directory",
        lambda: tmp_path,
    )
    fixed = datetime.datetime(2026, 10, 22, 12, 0, 0)

    class PatchedDatetime(datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            return fixed

    monkeypatch.setattr(
        "genesys_code_compare.file.sheet_service.datetime", PatchedDatetime
    )
    path = service.get_output_path(None)

    assert path == tmp_path / "reports" / "report_20261022-120000.xlsx"


def test_given_path(service, tmp_path):
    path = service.get_output_path(tmp_path / "reports" / "activity_codes.xlsx")
    assert path == tmp_path / "reports" / "activity_codes.xlsx"
