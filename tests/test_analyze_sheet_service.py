import pytest
from openpyxl import load_workbook

from genesys_code_compare.cli import CLI
from genesys_code_compare.file.analyze_sheet_service import AnalyzeSheetService
from genesys_code_compare.models.analyze import CodePresence, MatchResult
from genesys_code_compare.models.genesys import ActivityCode, BusinessUnit


@pytest.fixture
def service():
    return AnalyzeSheetService(CLI())


@pytest.fixture
def business_units():
    # Deliberately use a different order from the presence dictionaries.
    return [
        BusinessUnit(bu_id="b", name="BU B"),
        BusinessUnit(bu_id="a", name="BU A"),
        BusinessUnit(bu_id="c", name="BU C"),
    ]


def make_presence(name, sources, categories=None):
    categories = categories or {}
    return CodePresence(
        name=name,
        present_in={
            source: ActivityCode(
                ac_id=f"{source}-{name}",
                name=name,
                default_code=True,
                category=categories.get(source, "Break"),
                default_name_changed=False,
            )
            for source in sources
        },
        missing_from=[
            source for source in ["BU A", "BU B", "BU C"] if source not in sources
        ],
    )


@pytest.fixture
def match_result():
    exact = make_presence("Lunch", ["BU A", "BU B", "BU C"])
    different = make_presence(
        "Training", ["BU A", "BU B", "BU C"], {"BU B": "Meeting"}
    )
    partial = make_presence("Meeting", ["BU A", "BU C"])
    unique = make_presence("Coaching", ["BU B"])
    return MatchResult(
        exact_name_matches=[exact, different],
        exact_field_matches=[exact],
        field_differences=[different],
        partial_matches=[partial],
        unique=[unique],
    )


def test_summary_headers_follow_business_unit_order(service, business_units):
    sheet = service._build_summary(business_units, MatchResult())

    assert sheet.headers == ["Code", "BU B", "BU A", "BU C"]


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("Lunch", ["X", "X", "X"]),
        ("Training", ["X", "X", "X"]),
        ("Meeting", ["O", "X", "X"]),
        ("Coaching", ["X", "O", "O"]),
    ],
)
def test_summary_presence_aligns_with_headers(
    service, business_units, match_result, name, expected
):
    sheet = service._build_summary(business_units, match_result)

    row = next(row for row in sheet.rows if row[0] == name)
    assert row == [name, *expected]


def test_summary_includes_each_code_once(service, business_units, match_result):
    sheet = service._build_summary(business_units, match_result)

    assert sorted(row[0] for row in sheet.rows) == [
        "Coaching", "Lunch", "Meeting", "Training"
    ]


def test_summary_without_matches_has_no_rows(service, business_units):
    sheet = service._build_summary(business_units, MatchResult())

    assert list(sheet.rows) == []


def test_detail_uses_requested_title(service):
    sheet = service._build_detail_sheet("Unique Codes", [])

    assert sheet.name == "Unique Codes"


def test_detail_headers(service):
    sheet = service._build_detail_sheet("Details", [])

    assert sheet.headers == [
        "Code", "Business Unit", "Category", "Default Code",
        "Default Name Changed", "ID",
    ]


def test_detail_preserves_activity_fields(service):
    code = ActivityCode("id-42", "Training", False, "Meeting", True)
    record = CodePresence("Training", {"BU A": code}, ["BU B"])

    sheet = service._build_detail_sheet("Details", [record])

    assert list(sheet.rows) == [
        ["Training", "BU A", "Meeting", False, True, "id-42"]
    ]


def test_detail_sorts_codes_then_business_units(service):
    records = [
        make_presence("Training", ["BU B", "BU A"]),
        make_presence("Lunch", ["BU C", "BU A"]),
    ]

    sheet = service._build_detail_sheet("Details", records)

    assert [row[:2] for row in sheet.rows] == [
        ["Lunch", "BU A"], ["Lunch", "BU C"],
        ["Training", "BU A"], ["Training", "BU B"],
    ]


def test_detail_without_matches_has_no_rows(service):
    sheet = service._build_detail_sheet("Details", [])

    assert list(sheet.rows) == []


@pytest.fixture
def saved_workbook(service, business_units, match_result, tmp_path):
    output = tmp_path / "reports" / "analysis.xlsx"
    service.write_analyze(business_units, match_result, output)
    workbook = load_workbook(output)
    yield workbook
    workbook.close()


def test_write_creates_all_category_tabs(saved_workbook):
    assert saved_workbook.sheetnames == [
        "Summary", "Exact Matches", "Field Differences",
        "Partial Matches", "Unique Codes",
    ]


@pytest.mark.parametrize(
    ("title", "expected_names"),
    [
        ("Summary", ["Lunch", "Training", "Meeting", "Coaching"]),
        ("Exact Matches", ["Lunch", "Lunch", "Lunch"]),
        ("Field Differences", ["Training", "Training", "Training"]),
        ("Partial Matches", ["Meeting", "Meeting"]),
        ("Unique Codes", ["Coaching"]),
    ],
)
def test_write_routes_records_to_correct_tab(saved_workbook, title, expected_names):
    rows = saved_workbook[title].iter_rows(min_row=2, values_only=True)

    assert [row[0] for row in rows] == expected_names


def test_write_returns_output_path(service, business_units, tmp_path):
    output = tmp_path / "empty.xlsx"

    result = service.write_analyze(business_units, MatchResult(), output)

    assert result == output
