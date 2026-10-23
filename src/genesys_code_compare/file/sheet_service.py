from collections.abc import Iterable

from ..cli import CLI
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from ..models.file import SheetData
from .util import FileUtil
from datetime import datetime


class SheetService:

    def __init__(self, cli: CLI):
        self._cli = cli

    def write(self, output_path: str | Path, sheets: Iterable[SheetData]) -> Path:
        workbook = Workbook()
        workbook.remove(workbook.active)

        for data in sheets:
            sheet = workbook.create_sheet(title=data.name)
            sheet.append(list(data.headers))

            for row in data.rows:
                sheet.append(list(row))

            for cell in sheet[1]:
                cell.font = Font(bold=True)

            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"

            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions

        if not workbook.worksheets:
            raise ValueError("Provide at least one sheet.")

        output_path = self.get_output_path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        self._cli.print_status(f"Wrote file to: {output_path}")
        return output_path

    def get_output_path(self, output_path: str | Path | None = None) -> Path:
        if output_path is not None:
            return Path(output_path)
        now = datetime.now()
        timestamp_str = now.strftime("%Y%m%d-%H%M%S")
        return FileUtil.get_app_directory() / "reports" / f"report_{timestamp_str}.xlsx"
